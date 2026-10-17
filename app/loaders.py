"""Document loaders for PDF, Word (.docx) and Excel (.xlsx) files.

Each loader returns the plain text extracted from the file. Keeping extraction
simple and transparent matters here: the extracted text becomes the *only*
source of truth the assistant is allowed to use, so we want it to be faithful.
"""

from __future__ import annotations

import io
import os

from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".txt", ".md")


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a file given its name and raw bytes.

    Args:
        filename: Original filename (used only to detect the extension).
        data: Raw file bytes.

    Returns:
        Extracted text. Empty string if the format is unsupported.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext == ".xlsx":
        return _extract_xlsx(data)
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")

    return ""


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            # Tag the page so citations can point to it.
            pages.append(f"[page {i}]\n{text}")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables, which Word docs often use for specs.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
